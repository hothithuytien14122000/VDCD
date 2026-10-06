"""
Tao file Excel tong hop test case tu testcases_export.json (sinh boi export_json.js).
Moi NHOM API (Survey, SurveyForm, SurveyPublic, SurveyReport, SurveyResult) -> 1 sheet rieng.
Trong moi sheet, KHONG co cot Endpoint - thay vao do moi endpoint co 1 dong tieu de (banner)
rieng, ben duoi la cac test case cua chinh endpoint do. Co them 1 sheet "Index" tong hop dau tien.

Cot du lieu (dung thu tu theo yeu cau): ID, TITLE, DIEU KIEN TIEN QUYET, STEP,
KET QUA MONG DOI, DO UU TIEN, STATUS KHI RUN TEST (cot trong, co dropdown, de QA dien khi chay test).

Chay: python scripts/build_xlsx.py
"""
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
JSON_PATH = os.path.join(BASE_DIR, "testcases_export.json")
OUT_PATH = os.path.join(BASE_DIR, "testcases_export.xlsx")

FONT_NAME = "Arial"

TITLE_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True, color="1F4E78")
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, italic=True, color="666666")

ENDPOINT_FILL = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
ENDPOINT_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")

NORMAL_FONT = Font(name=FONT_NAME, size=10)
BOLD_FONT = Font(name=FONT_NAME, size=10, bold=True)
LINK_FONT = Font(name=FONT_NAME, size=10, color="1155CC", underline="single")

MUTATION_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
PRIORITY_COLORS = {
    "Cao": Font(name=FONT_NAME, size=10, bold=True, color="C00000"),
    "Trung binh": Font(name=FONT_NAME, size=10, color="B36B00"),
    "Thap": Font(name=FONT_NAME, size=10, color="1F7A1F"),
}

STATUS_OPTIONS = ["Dat", "Khong dat", "Chan (Blocked)", "Chua thuc hien"]
STATUS_FILL = {
    "Dat": PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid"),
    "Khong dat": PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
    "Chan (Blocked)": PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid"),
}

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WRAP_TOP = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_TOP = Alignment(wrap_text=True, vertical="top", horizontal="center")

# (tieu de cot, key trong JSON, do rong cot) - KHONG co Endpoint, dung thu tu theo yeu cau
DATA_COLUMNS = [
    ("ID", "id", 16),
    ("TITLE", "title", 40),
    ("DIEU KIEN TIEN QUYET", "preconditions", 36),
    ("STEP", "steps", 50),
    ("KET QUA MONG DOI", "expected_result", 50),
    ("DO UU TIEN", "priority", 12),
    ("STATUS KHI RUN TEST", None, 20),
]
N_COLS = len(DATA_COLUMNS)
LAST_COL_LETTER = get_column_letter(N_COLS)
STATUS_COL = N_COLS  # cot cuoi cung

GROUP_DESCRIPTIONS = {
    "Survey": "Quan ly dot khao sat: danh sach, tao, sua, xoa, phe duyet (UC-E.1, UC-E.2, UC-E.3)",
    "SurveyForm": "Tai mau phieu khao sat (.doc) de in phat cho nguoi dan",
    "SurveyPublic": "API cong khai cho nguoi dan thuc hien khao sat truc tuyen (UC-E.4)",
    "SurveyReport": "Bao cao, tong ket ket qua khao sat da hoan thanh (UC-E.5)",
    "SurveyResult": "Ghi nhan/cap nhat ket qua khao sat truc tiep, gian tiep qua phieu (UC-E.6)",
}


def count_lines(text):
    if not text:
        return 1
    return max(1, text.count("\n") + 1)


def group_by_endpoint(test_cases):
    """Gom cac test case lien tiep co cung endpoint thanh tung khoi (giu nguyen thu tu)."""
    blocks = []
    for tc in test_cases:
        ep = tc.get("endpoint", "")
        if blocks and blocks[-1]["endpoint"] == ep:
            blocks[-1]["items"].append(tc)
        else:
            blocks.append({"endpoint": ep, "items": [tc]})
    return blocks


def style_header_row(ws, row):
    for col_idx, (header, _, width) in enumerate(DATA_COLUMNS, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_TOP
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[row].height = 26


def build_group_sheet(wb, group):
    sheet_name = group["group"][:31]
    ws = wb.create_sheet(title=sheet_name)
    ws.sheet_view.showGridLines = False

    ws.merge_cells(f"A1:{LAST_COL_LETTER}1")
    ws["A1"] = f"Nhom API: {group['group']}"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(f"A2:{LAST_COL_LETTER}2")
    ws["A2"] = GROUP_DESCRIPTIONS.get(group["group"], group.get("module", ""))
    ws["A2"].font = SUBTITLE_FONT
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 16

    status_ranges = []  # danh sach (start_row, end_row) chua cot STATUS can ap dung dropdown
    r = 3  # dong trong (spacer) sau phan tieu de nhom

    for block in group_by_endpoint(group["test_cases"]):
        r += 1  # dong trong truoc moi endpoint
        banner_row = r
        ws.merge_cells(start_row=banner_row, start_column=1, end_row=banner_row, end_column=N_COLS)
        banner_cell = ws.cell(row=banner_row, column=1, value=f"Endpoint: {block['endpoint']}")
        banner_cell.font = ENDPOINT_FONT
        banner_cell.fill = ENDPOINT_FILL
        banner_cell.alignment = Alignment(vertical="center", horizontal="left")
        ws.row_dimensions[banner_row].height = 20
        r += 1

        header_row = r
        style_header_row(ws, header_row)
        r += 1

        data_start = r
        for tc in block["items"]:
            is_mutation = "mutation" in (tc.get("tags") or "")
            max_lines = 1
            for col_idx, (_, key, _) in enumerate(DATA_COLUMNS, start=1):
                if key is None:
                    cell = ws.cell(row=r, column=col_idx, value=None)
                else:
                    value = tc.get(key, "")
                    cell = ws.cell(row=r, column=col_idx, value=value)
                    max_lines = max(max_lines, count_lines(value))
                cell.border = BORDER
                cell.alignment = WRAP_TOP
                cell.font = NORMAL_FONT
                if col_idx == 1:
                    cell.font = BOLD_FONT
                if key == "priority" and cell.value in PRIORITY_COLORS:
                    cell.font = PRIORITY_COLORS[cell.value]
                if is_mutation:
                    cell.fill = MUTATION_FILL
            ws.row_dimensions[r].height = min(15 * max_lines, 300)
            r += 1
        data_end = r - 1
        if data_end >= data_start:
            status_ranges.append((data_start, data_end))

    # Dropdown + to mau co dieu kien cho cot STATUS KHI RUN TEST, chi ap dung cho cac dong du lieu
    if status_ranges:
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_OPTIONS) + '"', allow_blank=True)
        dv.error = "Vui long chon 1 trong cac gia tri co san"
        dv.errorTitle = "Gia tri khong hop le"
        ws.add_data_validation(dv)
        status_letter = get_column_letter(STATUS_COL)
        for start, end in status_ranges:
            dv.add(f"{status_letter}{start}:{status_letter}{end}")
            for status_value, fill in STATUS_FILL.items():
                ws.conditional_formatting.add(
                    f"{status_letter}{start}:{status_letter}{end}",
                    CellIsRule(operator="equal", formula=[f'"{status_value}"'], fill=fill),
                )

    ws.freeze_panes = "A3"
    return sum(len(b["items"]) for b in group_by_endpoint(group["test_cases"]))


def build_index_sheet(wb, groups_meta):
    ws = wb.create_sheet(title="Index", index=0)
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:D1")
    ws["A1"] = "TONG HOP TEST CASE API - PHAN HE KHAO SAT MUC DO HAI LONG CUA NGUOI DAN"
    ws["A1"].font = Font(name=FONT_NAME, size=15, bold=True, color="1F4E78")
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:D2")
    ws["A2"] = "Nguon: cac file testcases/**/*.yaml  |  Server: https://cds.geneat.pro  |  Moi nhom API la 1 sheet rieng"
    ws["A2"].font = SUBTITLE_FONT

    headers = ["STT", "Nhom API (Sheet)", "Mo ta", "So luong Test Case"]
    header_row = 4
    widths = [6, 22, 70, 20]
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=header_row, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_TOP
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[header_row].height = 20

    r = header_row + 1
    first_data_row = r
    for i, meta in enumerate(groups_meta, start=1):
        sheet_name = meta["sheetName"]
        ws.cell(row=r, column=1, value=i).alignment = CENTER_TOP
        link_cell = ws.cell(row=r, column=2, value=sheet_name)
        link_cell.hyperlink = f"#'{sheet_name}'!A1"
        link_cell.font = LINK_FONT
        ws.cell(row=r, column=3, value=GROUP_DESCRIPTIONS.get(sheet_name, ""))

        # Dem bang COUNTIF theo tien to "TC_" cua ID - khong phu thuoc vi tri dong (ro bust truoc
        # cac dong banner/endpoint va header lap lai trong sheet nhom)
        count_cell = ws.cell(row=r, column=4, value=f'=COUNTIF(\'{sheet_name}\'!A:A,"TC_*")')
        count_cell.alignment = CENTER_TOP

        for c in range(1, 5):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = NORMAL_FONT
            if c == 3:
                cell.alignment = WRAP_TOP
        r += 1

    total_row = r
    ws.cell(row=total_row, column=2, value="TONG CONG").font = BOLD_FONT
    total_cell = ws.cell(row=total_row, column=4, value=f"=SUM(D{first_data_row}:D{r - 1})")
    total_cell.font = BOLD_FONT
    for c in range(1, 5):
        ws.cell(row=total_row, column=c).border = BORDER

    legend_row = total_row + 2
    ws.cell(row=legend_row, column=1, value="Chu thich:").font = BOLD_FONT
    legends = [
        "Moi NHOM API (tag Swagger) la 1 sheet rieng. Trong moi sheet, moi ENDPOINT co 1 dong tieu de (banner) mau xanh,"
        " ben duoi la cac test case cua chinh endpoint do (khong co cot Endpoint rieng).",
        "Hang co nen vang nhat = test case [Mutation] (tao/sua/xoa du lieu that tren server).",
        "Do uu tien: Cao (chu do), Trung binh (chu cam), Thap (chu xanh).",
        "Cot 'STATUS KHI RUN TEST' co san dropdown (Dat / Khong dat / Chan / Chua thuc hien) de QA dien khi thuc thi,"
        " tu dong to mau theo ket qua chon.",
        "Noi dung day du, chi tiet hon xem tai cac file testcases/<Nhom>/*.yaml tuong ung (khong thay doi).",
    ]
    for i, text in enumerate(legends, start=1):
        ws.cell(row=legend_row + i, column=1, value="- " + text).font = SUBTITLE_FONT
        ws.merge_cells(start_row=legend_row + i, start_column=1, end_row=legend_row + i, end_column=4)
        ws.row_dimensions[legend_row + i].height = 28

    ws.freeze_panes = f"A{header_row + 1}"


def main():
    with open(JSON_PATH, "r", encoding="utf-8") as f:
        data = json.load(f)

    wb = Workbook()
    wb.remove(wb.active)

    groups_meta = []
    for group in data["groups"]:
        n = build_group_sheet(wb, group)
        groups_meta.append({"sheetName": group["group"][:31], "n_test_cases": n})

    build_index_sheet(wb, groups_meta)
    wb.active = 0

    wb.save(OUT_PATH)
    total = sum(m["n_test_cases"] for m in groups_meta)
    msg = f"Da tao {OUT_PATH} voi {len(groups_meta) + 1} sheet (1 Index + {len(groups_meta)} nhom), tong {total} test case"
    print(msg.encode("ascii", "backslashreplace").decode("ascii"))


if __name__ == "__main__":
    main()
