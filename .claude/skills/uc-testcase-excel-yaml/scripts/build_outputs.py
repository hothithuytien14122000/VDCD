#!/usr/bin/env python3
"""Validate a test-case YAML, print coverage report, and render the 5-column Excel."""
import argparse, re, sys
from collections import defaultdict

import yaml

COLS = ["ID", "Title", "Điều kiện tiên quyết", "Step", "Kết quả mong đợi"]
TYPES = {"positive", "negative", "boundary", "permission", "state_transition",
         "decision_table", "error_handling", "ui"}
ACTIONS = {"navigate", "click", "fill", "select", "check", "upload", "wait", "press", "verify", "api_call"}
PRIORITIES = {"high", "medium", "low"}
AUTOMATABLE = {"yes", "partial", "no"}
ID_RE = re.compile(r"^TC_[A-Za-z0-9.]+_\d{3,}$")


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def validate(doc):
    errs, warns = [], []
    for k in ("schema_version", "meta", "requirements", "test_cases"):
        if k not in doc:
            errs.append(f"thiếu khóa gốc '{k}'")
    if errs:
        return errs, warns, {}
    reqs = {r.get("id"): r for r in doc["requirements"] or []}
    locators = (doc.get("automation_config") or {}).get("locators") or {}
    seen, sigs, covered = set(), {}, defaultdict(list)
    for tc in doc["test_cases"] or []:
        i = tc.get("id", "<no id>")
        e = lambda m: errs.append(f"{i}: {m}")
        w = lambda m: warns.append(f"{i}: {m}")
        if not ID_RE.match(str(i)): e("ID sai định dạng TC_<UC>_<NNN>")
        if i in seen: e("ID trùng")
        seen.add(i)
        for f in ("title", "expected_result"):
            if not tc.get(f): e(f"thiếu '{f}'")
        if not tc.get("preconditions"): e("thiếu 'preconditions'")
        if tc.get("type") not in TYPES: e(f"type không hợp lệ: {tc.get('type')}")
        if tc.get("priority") not in PRIORITIES: e("priority không hợp lệ")
        if (tc.get("automation") or {}).get("automatable") not in AUTOMATABLE: e("automation.automatable không hợp lệ")
        refs = tc.get("requirement_refs") or []
        if not refs: e("thiếu requirement_refs")
        for r in refs:
            if r not in reqs: e(f"requirement_refs '{r}' không tồn tại")
            covered[r].append(i)
        steps = tc.get("steps") or []
        if not steps: e("không có steps")
        for n, s in enumerate(steps, 1):
            if s.get("seq") != n: e(f"step seq phải liên tục, gặp {s.get('seq')} tại vị trí {n}")
            if s.get("action") not in ACTIONS: e(f"step {n}: action không hợp lệ '{s.get('action')}'")
            if not s.get("description"): e(f"step {n}: thiếu description")
            ex = s.get("expected")
            if not isinstance(ex, dict) or "description" not in ex: e(f"step {n}: thiếu expected.description")
            tgt = s.get("target")
            if tgt and tgt not in locators:
                w(f"step {n}: target '{tgt}' chưa khai báo trong automation_config.locators")
            if "${" in str(s.get("input") or ""):
                for v in re.findall(r"\$\{(\w+)\}", s["input"]):
                    if v not in (tc.get("test_data") or {}): e(f"step {n}: biến ${{{v}}} không có trong test_data")
        sig = (norm(tc.get("title")), tuple(norm(s.get("description")) + norm(s.get("input")) for s in steps))
        if sig in sigs: w(f"trùng lặp với {sigs[sig]}")
        sigs[sig] = i
    return errs, warns, {"reqs": reqs, "covered": covered}


def render_xlsx(doc, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    wb = Workbook(); ws = wb.active; ws.title = "Test Cases"
    ws.append(COLS)
    for tc in doc["test_cases"]:
        pre = "\n".join(f"- {p}" for p in tc["preconditions"])
        st = "\n".join(f"{s['seq']}. {s['description']}" for s in tc["steps"])
        ex = [f"{s['seq']}. {s['expected']['description']}" for s in tc["steps"]
              if norm(s["expected"]["description"]) not in ("", "—", "-")]
        ex.append(f"Kết quả cuối: {tc['expected_result']}")
        ws.append([tc["id"], tc["title"], pre, st, "\n".join(ex)])
    thin = Side(style="thin", color="999999")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for col, wd in zip("ABCDE", (18, 40, 40, 60, 60)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("yaml_file"); ap.add_argument("--xlsx", help="đường dẫn Excel đầu ra (thư mục TC-xlsx/)")
    a = ap.parse_args()
    with open(a.yaml_file, encoding="utf-8") as f:
        doc = yaml.safe_load(f)
    errs, warns, info = validate(doc)
    for m in errs: print("LỖI   :", m)
    for m in warns: print("CẢNH BÁO:", m)
    if errs:
        print(f"\n{len(errs)} lỗi — chưa xuất Excel."); sys.exit(1)
    tcs = doc["test_cases"]
    print(f"\nTổng TC: {len(tcs)}")
    by = defaultdict(int)
    for t in tcs: by[t["type"]] += 1
    print("Theo loại:", dict(by))
    unc = [r for r in info["reqs"] if r not in info["covered"]]
    print("Requirement chưa có TC:", unc or "không")
    oq = doc.get("open_questions") or []
    print(f"Câu hỏi chưa rõ: {len(oq)}")
    for q in oq: print(f"  - {q.get('id')}: {q.get('text')}")
    blocked = [(t["id"], t["automation"].get("blockers")) for t in tcs if t["automation"]["automatable"] != "yes"]
    print(f"TC chưa automate được đầy đủ: {len(blocked)}/{len(tcs)}")
    missing = sorted({k for k, v in ((doc.get('automation_config') or {}).get('locators') or {}).items()
                      if (v or {}).get('status') != 'confirmed'})
    print("Locator thiếu/chưa xác nhận:", missing or "không")
    if a.xlsx:
        import os; os.makedirs(os.path.dirname(os.path.abspath(a.xlsx)), exist_ok=True)
        render_xlsx(doc, a.xlsx); print("Đã xuất:", a.xlsx)


if __name__ == "__main__":
    main()
