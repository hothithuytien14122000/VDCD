#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
render_xlsx.py — Render one or more test-case JSON files (see
reference/testcase-format.md) into a styled .xlsx, reusing the exact look already
used across testcases-xlsx/*.xlsx in this repo (title/subtitle banners, blue
"Chuc nang: ..." function banner, header row, priority colors, STATUS dropdown with
conditional formatting, edge/bug row highlighting, and a hyperlinked Index sheet).

Usage:
    python3 render_xlsx.py --in blocks.json --sheet-name Dang-Nhap \
        --out testcases-xlsx/uc1-1-dang-nhap-he-thong-test-cases.xlsx

    # Combine several UCs into one workbook (repeat --in/--sheet-name in order):
    python3 render_xlsx.py \
        --in uc1_1.json --sheet-name DangNhap \
        --in uc1_2.json --sheet-name DoiMatKhau \
        --out testcases-xlsx/combined-test-cases.xlsx

Input JSON shape:
    {
      "group_title": "...",
      "group_subtitle": "...",
      "blocks": [{"function": "...", "items": [{"id","title","pre","steps","expected","priority"}, ...]}]
    }
"""
import argparse
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

if sys.platform == 'win32':
    import codecs
    sys.stdout = codecs.getwriter('utf-8')(sys.stdout.buffer, 'strict')
    sys.stderr = codecs.getwriter('utf-8')(sys.stderr.buffer, 'strict')

FONT_NAME = "Arial"

TITLE_FONT = Font(name=FONT_NAME, size=15, bold=True, color="1F4E78")
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, italic=True, color="666666")

FUNC_FILL = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
FUNC_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")

NORMAL_FONT = Font(name=FONT_NAME, size=10)
BOLD_FONT = Font(name=FONT_NAME, size=10, bold=True)
LINK_FONT = Font(name=FONT_NAME, size=10, color="1155CC", underline="single")

EDGE_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
BUG_FILL = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")

PRIORITY_COLORS = {
    "Cao": Font(name=FONT_NAME, size=10, bold=True, color="C00000"),
    "Trung binh": Font(name=FONT_NAME, size=10, color="B36B00"),
    "Thap": Font(name=FONT_NAME, size=10, color="1F7A1F"),
}

STATUS_OPTIONS = ["Dat", "Khong dat", "Chan (Blocked)", "Chua thuc hien"]
STATUS_FILL = {
    "Dat": PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid"),
    "Khong dat": PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
    "Chan (Blocked)": PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid"),
}

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WRAP_TOP = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_TOP = Alignment(wrap_text=True, vertical="top", horizontal="center")

DATA_COLUMNS = [
    ("ID", "id", 16),
    ("TITLE", "title", 42),
    ("DIEU KIEN TIEN QUYET", "pre", 34),
    ("STEP", "steps", 46),
    ("KET QUA MONG DOI", "expected", 54),
    ("DO UU TIEN", "priority", 12),
    ("STATUS KHI RUN TEST", None, 20),
]
N_COLS = len(DATA_COLUMNS)
LAST_COL_LETTER = get_column_letter(N_COLS)
STATUS_COL = N_COLS


def count_lines(text):
    if not text:
        return 1
    return max(1, str(text).count("\n") + 1)


def id_sort_key(tc):
    """Sort by the numeric tail of the id (e.g. TC_UC2.1_009), with an optional
    trailing letter suffix (TC_UC2.1_025b) ordered right after its plain number
    (025, 025b, 025c, 026, ...). Falls back to the raw string for anything that
    doesn't match, so an unexpected id format never crashes the render."""
    m = re.search(r"(\d+)([a-zA-Z]*)$", str(tc.get("id", "")))
    if not m:
        return (1, str(tc.get("id", "")), "")
    return (0, int(m.group(1)), m.group(2))


def sorted_items(items):
    return sorted(items, key=id_sort_key)


def style_header_row(ws, row):
    for col_idx, (header, _, width) in enumerate(DATA_COLUMNS, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_TOP
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[row].height = 26


def build_main_sheet(wb, sheet_name, group_title, group_subtitle, blocks):
    ws = wb.create_sheet(title=sheet_name)
    ws.sheet_view.showGridLines = False

    ws.merge_cells(f"A1:{LAST_COL_LETTER}1")
    ws["A1"] = group_title
    ws["A1"].font = TITLE_FONT
    ws.row_dimensions[1].height = 24

    ws.merge_cells(f"A2:{LAST_COL_LETTER}2")
    ws["A2"] = group_subtitle
    ws["A2"].font = SUBTITLE_FONT
    ws.row_dimensions[2].height = 16

    status_ranges = []
    r = 3
    total_tc = 0
    toc = []  # (function_name, banner_row, n_items)

    for block in blocks:
        r += 1
        banner_row = r
        ws.merge_cells(start_row=banner_row, start_column=1, end_row=banner_row, end_column=N_COLS)
        banner_cell = ws.cell(row=banner_row, column=1, value=f"Chuc nang: {block['function']}")
        banner_cell.font = FUNC_FONT
        banner_cell.fill = FUNC_FILL
        banner_cell.alignment = Alignment(vertical="center", horizontal="left")
        ws.row_dimensions[banner_row].height = 20
        r += 1

        header_row = r
        style_header_row(ws, header_row)
        r += 1

        data_start = r
        for tc in sorted_items(block["items"]):
            title = tc.get("title", "")
            is_edge = "[Edge case]" in title or "[Validation]" in title
            is_bug = "[Bug ghi nhan]" in title
            max_lines = 1
            for col_idx, (_, key, _) in enumerate(DATA_COLUMNS, start=1):
                if key is None:
                    cell = ws.cell(row=r, column=col_idx, value=None)
                else:
                    value = tc.get(key, "")
                    cell = ws.cell(row=r, column=col_idx, value=value)
                    max_lines = max(max_lines, count_lines(value))
                cell.border = BORDER
                cell.alignment = WRAP_TOP
                cell.font = NORMAL_FONT
                if col_idx == 1:
                    cell.font = BOLD_FONT
                if key == "priority" and cell.value in PRIORITY_COLORS:
                    cell.font = PRIORITY_COLORS[cell.value]
                if is_bug:
                    cell.fill = BUG_FILL
                elif is_edge:
                    cell.fill = EDGE_FILL
            ws.row_dimensions[r].height = min(15 * max_lines, 320)
            r += 1
        data_end = r - 1
        if data_end >= data_start:
            status_ranges.append((data_start, data_end))
        n_items = len(block["items"])
        total_tc += n_items
        toc.append((block["function"], banner_row, n_items))
        r += 1  # blank row between blocks

    if status_ranges:
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_OPTIONS) + '"', allow_blank=True)
        dv.error = "Vui long chon 1 trong cac gia tri co san"
        dv.errorTitle = "Gia tri khong hop le"
        ws.add_data_validation(dv)
        status_letter = get_column_letter(STATUS_COL)
        for start, end in status_ranges:
            dv.add(f"{status_letter}{start}:{status_letter}{end}")
            for status_value, fill in STATUS_FILL.items():
                ws.conditional_formatting.add(
                    f"{status_letter}{start}:{status_letter}{end}",
                    CellIsRule(operator="equal", formula=[f'"{status_value}"'], fill=fill),
                )

    ws.freeze_panes = "A4"
    return total_tc, toc


def build_index_sheet(wb, sections):
    ws = wb.create_sheet(title="Index", index=0)
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:D1")
    ws["A1"] = "DANH SACH TEST CASE"
    ws["A1"].font = TITLE_FONT
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:D2")
    ws["A2"] = "Sinh boi uc-ui-testcase-generator - xem tung sheet du lieu de biet nguon UC/UI"
    ws["A2"].font = SUBTITLE_FONT

    headers = ["STT", "Chuc nang", "Vi tri", "So luong Test Case"]
    widths = [6, 62, 16, 20]

    r = 4
    grand_total_formula_parts = []
    for section in sections:
        sheet_name = section["sheet_name"]
        toc = section["toc"]

        section_title_row = r
        ws.merge_cells(start_row=section_title_row, start_column=1, end_row=section_title_row, end_column=4)
        title_cell = ws.cell(row=section_title_row, column=1, value=f"Sheet: {sheet_name}  -  {section['group_title']}")
        title_cell.font = FUNC_FONT
        title_cell.fill = FUNC_FILL
        title_cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        ws.row_dimensions[section_title_row].height = 20
        r += 1

        header_row = r
        for i, (h, w) in enumerate(zip(headers, widths), start=1):
            cell = ws.cell(row=header_row, column=i, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = CENTER_TOP
            cell.border = BORDER
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[header_row].height = 20
        r += 1

        first_data_row = r
        for i, (func_name, banner_row, n_items) in enumerate(toc, start=1):
            ws.cell(row=r, column=1, value=i).alignment = CENTER_TOP
            link_cell = ws.cell(row=r, column=2, value=func_name)
            link_cell.hyperlink = f"#'{sheet_name}'!A{banner_row}"
            link_cell.font = LINK_FONT
            pos_cell = ws.cell(row=r, column=3, value=f"{sheet_name} - Dong {banner_row}")
            pos_cell.alignment = CENTER_TOP
            cnt_cell = ws.cell(row=r, column=4, value=n_items)
            cnt_cell.alignment = CENTER_TOP
            for c in range(1, 5):
                cell = ws.cell(row=r, column=c)
                cell.border = BORDER
                cell.font = NORMAL_FONT
            r += 1

        subtotal_row = r
        ws.cell(row=subtotal_row, column=2, value=f"Cong ({sheet_name})").font = BOLD_FONT
        subtotal_cell = ws.cell(row=subtotal_row, column=4, value=f"=SUM(D{first_data_row}:D{r - 1})")
        subtotal_cell.font = BOLD_FONT
        for c in range(1, 5):
            ws.cell(row=subtotal_row, column=c).border = BORDER
        grand_total_formula_parts.append(f"D{subtotal_row}")
        r += 2

    grand_total_row = r
    ws.cell(row=grand_total_row, column=2, value="TONG CONG TOAN BO FILE").font = BOLD_FONT
    grand_total_cell = ws.cell(row=grand_total_row, column=4, value="=" + "+".join(grand_total_formula_parts))
    grand_total_cell.font = BOLD_FONT
    for c in range(1, 5):
        ws.cell(row=grand_total_row, column=c).border = BORDER

    ws.freeze_panes = "A5"


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--in", dest="inputs", action="append", required=True,
                     help="Path to a test-case JSON file (repeatable)")
    ap.add_argument("--sheet-name", dest="sheet_names", action="append", required=True,
                     help="Sheet name for the matching --in, same order (repeatable)")
    ap.add_argument("--group-title", dest="group_titles", action="append", default=[],
                     help="Override group_title for the matching --in, same order (optional)")
    ap.add_argument("--out", required=True, help="Output .xlsx path")
    args = ap.parse_args()

    if len(args.inputs) != len(args.sheet_names):
        print("ERROR: --in and --sheet-name must be passed the same number of times, in matching order.")
        sys.exit(1)

    wb = Workbook()
    wb.remove(wb.active)

    sections = []
    for i, (in_path, sheet_name) in enumerate(zip(args.inputs, args.sheet_names)):
        with open(in_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        group_title = data.get("group_title", sheet_name)
        if i < len(args.group_titles) and args.group_titles[i]:
            group_title = args.group_titles[i]
        group_subtitle = data.get("group_subtitle", "")
        blocks = data.get("blocks", [])

        total_tc, toc = build_main_sheet(wb, sheet_name, group_title, group_subtitle, blocks)
        sections.append(dict(sheet_name=sheet_name, group_title=group_title, total_tc=total_tc, toc=toc))

    build_index_sheet(wb, sections)
    wb.active = 0

    out_dir = os.path.dirname(args.out)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(args.out)

    grand_total = sum(s["total_tc"] for s in sections)
    sheet_names = ", ".join(s["sheet_name"] for s in sections)
    msg = f"Da tao {args.out} voi {len(sections) + 1} sheet (Index, {sheet_names}), tong {grand_total} test case"
    print(msg.encode("ascii", "backslashreplace").decode("ascii"))


if __name__ == "__main__":
    main()
