# -*- coding: utf-8 -*-
"""
Tao file Excel tong hop test case UI cho man hinh "Quan ly tu danh gia"
(Bo chi so DTI > Quan ly tu danh gia), theo dung format cot cua file mau
survey-api-tests/testcases_export.xlsx (ID, TITLE, DIEU KIEN TIEN QUYET, STEP,
KET QUA MONG DOI, DO UU TIEN, STATUS KHI RUN TEST), nhung banner nhom la
"Chuc nang: ..." thay vi "Endpoint: ..." vi day la test case UI, khong phai API.

Chay: python build_xlsx.py
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

from testcases_data import GROUP_TITLE, GROUP_SUBTITLE, BLOCKS

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(BASE_DIR, "testcases_quan-ly-tu-danh-gia.xlsx")

FONT_NAME = "Arial"

TITLE_FONT = Font(name=FONT_NAME, size=15, bold=True, color="1F4E78")
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, italic=True, color="666666")

FUNC_FILL = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
FUNC_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")

NORMAL_FONT = Font(name=FONT_NAME, size=10)
BOLD_FONT = Font(name=FONT_NAME, size=10, bold=True)
LINK_FONT = Font(name=FONT_NAME, size=10, color="1155CC", underline="single")

EDGE_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
BUG_FILL = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")

PRIORITY_COLORS = {
    "Cao": Font(name=FONT_NAME, size=10, bold=True, color="C00000"),
    "Trung binh": Font(name=FONT_NAME, size=10, color="B36B00"),
    "Thap": Font(name=FONT_NAME, size=10, color="1F7A1F"),
}

STATUS_OPTIONS = ["Dat", "Khong dat", "Chan (Blocked)", "Chua thuc hien"]
STATUS_FILL = {
    "Dat": PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid"),
    "Khong dat": PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
    "Chan (Blocked)": PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid"),
}

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WRAP_TOP = Alignment(wrap_text=True, vertical="top", horizontal="left")
CENTER_TOP = Alignment(wrap_text=True, vertical="top", horizontal="center")

# (tieu de cot, key trong dict test case, do rong cot)
DATA_COLUMNS = [
    ("ID", "id", 16),
    ("TITLE", "title", 42),
    ("DIEU KIEN TIEN QUYET", "pre", 34),
    ("STEP", "steps", 46),
    ("KET QUA MONG DOI", "expected", 54),
    ("DO UU TIEN", "priority", 12),
    ("STATUS KHI RUN TEST", None, 20),
]
N_COLS = len(DATA_COLUMNS)
LAST_COL_LETTER = get_column_letter(N_COLS)
STATUS_COL = N_COLS


def count_lines(text):
    if not text:
        return 1
    return max(1, text.count("\n") + 1)


def style_header_row(ws, row):
    for col_idx, (header, _, width) in enumerate(DATA_COLUMNS, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_TOP
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[row].height = 26


def build_main_sheet(wb):
    ws = wb.create_sheet(title="QuanLyTuDanhGia")
    ws.sheet_view.showGridLines = False

    ws.merge_cells(f"A1:{LAST_COL_LETTER}1")
    ws["A1"] = GROUP_TITLE
    ws["A1"].font = TITLE_FONT
    ws.row_dimensions[1].height = 24

    ws.merge_cells(f"A2:{LAST_COL_LETTER}2")
    ws["A2"] = GROUP_SUBTITLE
    ws["A2"].font = SUBTITLE_FONT
    ws.row_dimensions[2].height = 16

    status_ranges = []
    r = 3
    total_tc = 0
    toc = []  # (function_name, banner_row, n_items)

    for block in BLOCKS:
        r += 1
        banner_row = r
        ws.merge_cells(start_row=banner_row, start_column=1, end_row=banner_row, end_column=N_COLS)
        banner_cell = ws.cell(row=banner_row, column=1, value=f"Chuc nang: {block['function']}")
        banner_cell.font = FUNC_FONT
        banner_cell.fill = FUNC_FILL
        banner_cell.alignment = Alignment(vertical="center", horizontal="left")
        ws.row_dimensions[banner_row].height = 20
        r += 1

        header_row = r
        style_header_row(ws, header_row)
        r += 1

        data_start = r
        for tc in block["items"]:
            is_edge = "[Edge case]" in tc["title"] or "[Validation]" in tc["title"]
            is_bug = "[Bug ghi nhan]" in tc["title"]
            max_lines = 1
            for col_idx, (_, key, _) in enumerate(DATA_COLUMNS, start=1):
                if key is None:
                    cell = ws.cell(row=r, column=col_idx, value=None)
                else:
                    value = tc.get(key, "")
                    cell = ws.cell(row=r, column=col_idx, value=value)
                    max_lines = max(max_lines, count_lines(value))
                cell.border = BORDER
                cell.alignment = WRAP_TOP
                cell.font = NORMAL_FONT
                if col_idx == 1:
                    cell.font = BOLD_FONT
                if key == "priority" and cell.value in PRIORITY_COLORS:
                    cell.font = PRIORITY_COLORS[cell.value]
                if is_bug:
                    cell.fill = BUG_FILL
                elif is_edge:
                    cell.fill = EDGE_FILL
            ws.row_dimensions[r].height = min(15 * max_lines, 320)
            r += 1
        data_end = r - 1
        if data_end >= data_start:
            status_ranges.append((data_start, data_end))
        n_items = len(block["items"])
        total_tc += n_items
        toc.append((block["function"], banner_row, n_items))
        r += 1  # dong trong giua cac block

    if status_ranges:
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_OPTIONS) + '"', allow_blank=True)
        dv.error = "Vui long chon 1 trong cac gia tri co san"
        dv.errorTitle = "Gia tri khong hop le"
        ws.add_data_validation(dv)
        status_letter = get_column_letter(STATUS_COL)
        for start, end in status_ranges:
            dv.add(f"{status_letter}{start}:{status_letter}{end}")
            for status_value, fill in STATUS_FILL.items():
                ws.conditional_formatting.add(
                    f"{status_letter}{start}:{status_letter}{end}",
                    CellIsRule(operator="equal", formula=[f'"{status_value}"'], fill=fill),
                )

    ws.freeze_panes = "A4"
    return total_tc, toc


def build_index_sheet(wb, total_tc, toc):
    ws = wb.create_sheet(title="Index", index=0)
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:D1")
    ws["A1"] = GROUP_TITLE
    ws["A1"].font = TITLE_FONT
    ws.row_dimensions[1].height = 24

    ws.merge_cells("A2:D2")
    ws["A2"] = GROUP_SUBTITLE
    ws["A2"].font = SUBTITLE_FONT

    headers = ["STT", "Chuc nang (trong sheet QuanLyTuDanhGia)", "Vi tri", "So luong Test Case"]
    header_row = 4
    widths = [6, 62, 12, 20]
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=header_row, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_TOP
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[header_row].height = 20

    r = header_row + 1
    first_data_row = r
    for i, (func_name, banner_row, n_items) in enumerate(toc, start=1):
        ws.cell(row=r, column=1, value=i).alignment = CENTER_TOP
        link_cell = ws.cell(row=r, column=2, value=func_name)
        link_cell.hyperlink = f"#'QuanLyTuDanhGia'!A{banner_row}"
        link_cell.font = LINK_FONT
        pos_cell = ws.cell(row=r, column=3, value=f"Dong {banner_row}")
        pos_cell.alignment = CENTER_TOP
        cnt_cell = ws.cell(row=r, column=4, value=n_items)
        cnt_cell.alignment = CENTER_TOP
        for c in range(1, 5):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = NORMAL_FONT
        r += 1

    total_row = r
    ws.cell(row=total_row, column=2, value="TONG CONG").font = BOLD_FONT
    total_cell = ws.cell(row=total_row, column=4, value=f"=SUM(D{first_data_row}:D{r - 1})")
    total_cell.font = BOLD_FONT
    for c in range(1, 5):
        ws.cell(row=total_row, column=c).border = BORDER

    legend_row = total_row + 2
    ws.cell(row=legend_row, column=1, value="Chu thich:").font = BOLD_FONT
    legends = [
        "Sheet 'QuanLyTuDanhGia' gom tat ca test case UI cua menu Quan ly tu danh gia"
        " (Bo chi so DTI > Quan ly tu danh gia). Moi CHUC NANG co 1 dong tieu de (banner) mau xanh.",
        "Hang nen vang nhat = test case [Validation] / [Edge case]. Hang nen cam nhat = [Bug ghi nhan]"
        " (van de phat hien qua khao sat UI tu dong, can QA xac nhan lai thu cong).",
        "Do uu tien: Cao (chu do), Trung binh (chu cam), Thap (chu xanh).",
        "Cot 'STATUS KHI RUN TEST' co san dropdown (Dat / Khong dat / Chan / Chua thuc hien) de QA dien khi"
        " thuc thi, tu dong to mau theo ket qua chon.",
        "Cac noi dung danh dau 'CAN XAC NHAN' trong cot KET QUA MONG DOI la gia dinh nghiep vu dua tren"
        " UC_1409_Requirement.md (nhom B.III), chua duoc xac nhan 100% qua khao sat UI tu dong - QA can"
        " doi chieu voi dev/BA khi thuc thi.",
        "Nguon UI: khao sat truc tiep man hinh that bang Playwright (tai khoan qa.superadmin@cds.local)"
        " ngay 2026-10-06 tren https://cds.geneat.pro/#/vi/self-assessment-campaigns.",
    ]
    for i, text in enumerate(legends, start=1):
        ws.cell(row=legend_row + i, column=1, value="- " + text).font = SUBTITLE_FONT
        ws.merge_cells(start_row=legend_row + i, start_column=1, end_row=legend_row + i, end_column=4)
        ws.row_dimensions[legend_row + i].height = 30

    ws.freeze_panes = f"A{header_row + 1}"


def main():
    wb = Workbook()
    wb.remove(wb.active)

    total_tc, toc = build_main_sheet(wb)
    build_index_sheet(wb, total_tc, toc)
    wb.active = 0

    wb.save(OUT_PATH)
    msg = f"Da tao {OUT_PATH} voi 2 sheet (Index + QuanLyTuDanhGia), tong {total_tc} test case"
    print(msg.encode("ascii", "backslashreplace").decode("ascii"))


if __name__ == "__main__":
    main()
